"""Génère des fichiers Word (.docx) et PDF à partir de texte Markdown."""
import io
import re
import unicodedata
from pathlib import Path

FONTS = Path(__file__).parent / "fonts"
AR_RE = re.compile(r"[\u0590-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFF]")  # arabe, hébreu... (écritures de droite à gauche)
# caractères que les polices PDF standard (cp1252) savent afficher
NEEDS_UNICODE_RE = re.compile(r"[^\x00-\xff\u2018\u2019\u201c\u201d\u2013\u2014\u2026\u20ac\u2022\u202f\u2011\u2192\u2713\u2714\u2212]")

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


# ---------- analyse du Markdown ----------
def parse_blocks(md: str) -> list[tuple]:
    """Retourne une liste de blocs : ("h", niveau, texte), ("p", texte), ("ul", [..]),
    ("ol", [..]), ("table", [[cellules]]), ("hr",)."""
    blocks: list[tuple] = []
    lines = md.replace("\r", "").split("\n")
    para: list[str] = []

    def flush():
        if para:
            blocks.append(("p", "\n".join(para)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            flush()
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            blocks.append(("h", min(level, 4), s[level:].strip()))
        elif re.fullmatch(r"[-*_]{3,}", s):
            flush()
            blocks.append(("hr",))
        elif re.match(r"^[-*•]\s+", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*[-*•]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*•]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ul", items))
            continue
        elif re.match(r"^\d+[.)]\s+", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ol", items))
            continue
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


def slug(md: str) -> str:
    """Nom de fichier tiré du premier titre du document."""
    m = re.search(r"^#+\s*(.+)$", md, re.MULTILINE) or re.search(r"\S.*", md)
    title = m.group(1) if m and m.groups() else (m.group(0) if m else "document")
    title = re.sub(r"[*_`]", "", title)
    ascii_ = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    return (re.sub(r"[^a-z0-9]+", "-", ascii_.lower()).strip("-")[:50]) or "document"


# ---------- Word ----------
def _docx_runs(par, text: str):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _docx_lines(par, text: str):
    for n, line in enumerate(text.split("\n")):
        if n:
            par.add_run().add_break()
        _docx_runs(par, line)


def make_docx(md: str, path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 13), ("Heading 4", 11)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x6E)

    for b in parse_blocks(md):
        kind = b[0]
        if kind == "h":
            h = doc.add_heading(level=b[1])
            _docx_runs(h, b[2])
        elif kind == "p":
            _docx_lines(doc.add_paragraph(), b[1])
        elif kind in ("ul", "ol"):
            style = "List Bullet" if kind == "ul" else "List Number"
            for item in b[1]:
                _docx_runs(doc.add_paragraph(style=style), item)
        elif kind == "table":
            rows = b[1]
            ncols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = t.cell(ri, ci)
                    par = cell.paragraphs[0]
                    _docx_runs(par, row[ci] if ci < len(row) else "")
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:fill"), "DCE6F5")
                        tcPr.append(shd)
            doc.add_paragraph()
        elif kind == "hr":
            par = doc.add_paragraph()
            pPr = par._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(k), v)
            borders.append(bottom)
            pPr.append(borders)
    # paragraphes en arabe/hébreu : sens de lecture de droite à gauche
    def all_paragraphs():
        yield from doc.paragraphs
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    after = {"adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
             "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
             "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"}
    for par in all_paragraphs():
        if AR_RE.search(par.text):
            pPr = par._p.get_or_add_pPr()
            if pPr.find(qn("w:bidi")) is None:
                bidi = OxmlElement("w:bidi")
                nxt = next((c for c in pPr if c.tag.split("}")[1] in after), None)
                if nxt is not None:
                    nxt.addprevious(bidi)
                else:
                    pPr.append(bidi)
    doc.save(path)


# ---------- PDF ----------
_PDF_MAP = str.maketrans(
    {" ": " ", " ": " ", "‑": "-", "→": "->", "•": "-", "✓": "v", "✔": "v"}
)


def _pdf_safe(s: str) -> str:
    # Polices PDF standard = jeu cp1252 (accents français, €, guillemets OK ; pas d'emojis)
    return s.translate(_PDF_MAP).encode("cp1252", "ignore").decode("cp1252")


def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _rtl_line(text: str) -> str:
    """Texte arabe/hébreu prêt pour reportlab : lettres reliées (reshaper) puis ordre visuel (bidi)."""
    import arabic_reshaper
    from bidi.algorithm import get_display

    plain = re.sub(r"\*\*|\*|`", "", text)
    return get_display(arabic_reshaper.reshape(plain))


def _pdf_inline(text: str, uni: bool = False, wrap: tuple | None = None) -> str:
    """Convertit le Markdown en balises reportlab. `uni` : polices Unicode (arabe, etc.) au lieu de cp1252.
    `wrap` = (police, taille, largeur) : pour l'arabe, le texte est coupé en lignes AVANT l'ordre visuel."""
    if uni and AR_RE.search(text):
        import arabic_reshaper
        from bidi.algorithm import get_display
        from reportlab.lib.utils import simpleSplit

        lines = []
        for line in text.split("\n"):
            shaped = arabic_reshaper.reshape(re.sub(r"\*\*|\*|`", "", line))
            parts = simpleSplit(shaped, wrap[0], wrap[1], wrap[2]) if wrap else [shaped]
            lines += [get_display(p) for p in parts]
        return "<br/>".join(_esc(x) for x in lines)
    fix = (lambda x: x) if uni else _pdf_safe
    out = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append("<b>" + _esc(fix(part[2:-2])) + "</b>")
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append("<font name='Courier'>" + _esc(fix(part[1:-1])) + "</font>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append("<i>" + _esc(fix(part[1:-1])) + "</i>")
        else:
            out.append(_esc(fix(part)))
    return "".join(out).replace("\n", "<br/>")


def make_pdf(md: str, path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    uni = bool(NEEDS_UNICODE_RE.search(md))  # arabe, cyrillique, etc. : polices TrueType nécessaires
    if uni:
        for name, file in (("DejaVu", "DejaVuSans.ttf"), ("DejaVu-Bold", "DejaVuSans-Bold.ttf")):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(FONTS / file)))
        pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu", boldItalic="DejaVu-Bold")
        regular, bold = "DejaVu", "DejaVu-Bold"
    else:
        regular, bold = "Helvetica", "Helvetica-Bold"

    navy = colors.HexColor("#1F3A6E")
    body = ParagraphStyle("body", fontName=regular, fontSize=10.5, leading=15, spaceAfter=7, alignment=TA_LEFT)
    heads = {
        1: ParagraphStyle("h1", parent=body, fontName=bold, fontSize=20, leading=25, textColor=navy, spaceBefore=4, spaceAfter=10),
        2: ParagraphStyle("h2", parent=body, fontName=bold, fontSize=14.5, leading=19, textColor=navy, spaceBefore=12, spaceAfter=6),
        3: ParagraphStyle("h3", parent=body, fontName=bold, fontSize=12, leading=16, textColor=navy, spaceBefore=8, spaceAfter=4),
        4: ParagraphStyle("h4", parent=body, fontName=bold, fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=3),
    }
    cell = ParagraphStyle("cell", parent=body, fontSize=9.5, leading=12, spaceAfter=0)
    cell_b = ParagraphStyle("cellb", parent=cell, fontName=bold)
    rtl_cache = {}

    FRAME_W = A4[0] - 4.4 * cm

    def P(text, style, width=None):
        """Paragraphe ; aligné à droite si le texte est en arabe/hébreu."""
        if uni and AR_RE.search(text):
            style = rtl_cache.setdefault(style.name, ParagraphStyle(style.name + "_r", parent=style, alignment=TA_RIGHT))
        return Paragraph(_pdf_inline(text, uni, (style.fontName, style.fontSize, (width or FRAME_W) - 2)), style)

    story = []
    for b in parse_blocks(md):
        kind = b[0]
        if kind == "h":
            story.append(P(b[2], heads[b[1]]))
        elif kind == "p":
            story.append(P(b[1], body))
        elif kind in ("ul", "ol"):
            if uni and any(AR_RE.search(t) for t in b[1]):
                for i, t in enumerate(b[1], 1):  # puce ou numéro en début de texte : affiché à droite
                    story.append(P(("•  " if kind == "ul" else f"{i}. ") + t, body, FRAME_W - 12))
            else:
                items = [ListItem(P(t, body), leftIndent=14) for t in b[1]]
                story.append(
                    ListFlowable(items, bulletType="bullet" if kind == "ul" else "1", start="-" if kind == "ul" else 1, leftIndent=16)
                )
            story.append(Spacer(1, 4))
        elif kind == "table":
            rows = b[1]
            ncols = max(len(r) for r in rows)
            colw = FRAME_W / ncols
            order = range(ncols - 1, -1, -1) if (uni and any(AR_RE.search(c) for c in rows[0])) else range(ncols)
            data = [
                [P(r[c] if c < len(r) else "", cell_b if ri == 0 else cell, colw - 12) for c in order]
                for ri, r in enumerate(rows)
            ]
            t = Table(data, repeatRows=1, colWidths=[colw] * ncols)
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F5")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story += [t, Spacer(1, 10)]
        elif kind == "hr":
            story.append(HRFlowable(width="100%", color=colors.HexColor("#999999"), spaceBefore=4, spaceAfter=8))

    SimpleDocTemplate(
        path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm
    ).build(story or [Paragraph("(document vide)", body)])


def render(fmt: str, md: str) -> tuple[str, bytes]:
    """Crée le document en mémoire (fmt = 'docx' ou 'pdf'). Retourne (nom de fichier, octets)."""
    buf = io.BytesIO()
    (make_docx if fmt == "docx" else make_pdf)(md, buf)
    return f"{slug(md)}.{fmt}", buf.getvalue()
